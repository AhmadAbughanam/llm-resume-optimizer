import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk
from core.parser import parse_file
from core import scorer
import threading
from docx import Document
from tkinter.filedialog import asksaveasfilename
import re
from core.rewriter import rewrite_sections_streaming


class ResumeOptimizerGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("LLM Resume Optimizer")
        self.root.geometry("1200x800")
        self.root.configure(bg="#f0f0f0")

        # Configure style
        self.setup_styles()
        self.create_widgets()

    def setup_styles(self):
        """Setup ttk styles for better appearance"""
        style = ttk.Style()
        style.theme_use("clam")

        # Configure custom styles
        style.configure(
            "Title.TLabel", font=("Arial", 16, "bold"), background="#f0f0f0"
        )
        style.configure(
            "Heading.TLabel", font=("Arial", 12, "bold"), background="#f0f0f0"
        )
        style.configure("Action.TButton", font=("Arial", 10, "bold"), padding=10)
        style.configure(
            "Score.TLabel",
            font=("Arial", 11),
            background="white",
            relief="solid",
            borderwidth=1,
            padding=10,
        )

    def create_widgets(self):
        # Main container
        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        # Configure grid weights
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(0, weight=1)
        main_frame.grid_rowconfigure(2, weight=1)
        main_frame.grid_columnconfigure(0, weight=1)
        main_frame.grid_columnconfigure(1, weight=1)

        # Title
        title_label = ttk.Label(
            main_frame, text="🚀 LLM Resume Optimizer", style="Title.TLabel"
        )
        title_label.grid(row=0, column=0, columnspan=2, pady=(0, 20))

        # File upload section
        upload_frame = ttk.LabelFrame(main_frame, text="📁 File Upload", padding="15")
        upload_frame.grid(
            row=1, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=(0, 15)
        )
        upload_frame.grid_columnconfigure(0, weight=1)
        upload_frame.grid_columnconfigure(1, weight=1)

        # Upload buttons
        ttk.Button(
            upload_frame,
            text="📄 Upload Resume",
            command=self.upload_resume,
            style="Action.TButton",
        ).grid(row=0, column=0, padx=(0, 10), sticky=(tk.W, tk.E))

        ttk.Button(
            upload_frame,
            text="📋 Upload Job Description",
            command=self.upload_job,
            style="Action.TButton",
        ).grid(row=0, column=1, padx=(10, 0), sticky=(tk.W, tk.E))

        # Content area
        content_frame = ttk.Frame(main_frame)
        content_frame.grid(
            row=2, column=0, columnspan=2, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(0, 15)
        )
        content_frame.grid_rowconfigure(1, weight=1)
        content_frame.grid_rowconfigure(3, weight=1)
        content_frame.grid_columnconfigure(0, weight=1)
        content_frame.grid_columnconfigure(1, weight=1)

        # Resume section
        resume_label = ttk.Label(
            content_frame, text="📝 Resume Content", style="Heading.TLabel"
        )
        resume_label.grid(row=0, column=0, sticky=tk.W, padx=(0, 10), pady=(0, 5))

        resume_frame = ttk.Frame(content_frame)
        resume_frame.grid(
            row=1, column=0, sticky=(tk.W, tk.E, tk.N, tk.S), padx=(0, 10)
        )
        resume_frame.grid_rowconfigure(0, weight=1)
        resume_frame.grid_columnconfigure(0, weight=1)

        self.resume_text = scrolledtext.ScrolledText(
            resume_frame,
            wrap=tk.WORD,
            font=("Consolas", 10),
            bg="white",
            relief="solid",
            borderwidth=1,
        )
        self.resume_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        # Job description section
        job_label = ttk.Label(
            content_frame, text="💼 Job Description Content", style="Heading.TLabel"
        )
        job_label.grid(row=0, column=1, sticky=tk.W, padx=(10, 0), pady=(0, 5))

        job_frame = ttk.Frame(content_frame)
        job_frame.grid(row=1, column=1, sticky=(tk.W, tk.E, tk.N, tk.S), padx=(10, 0))
        job_frame.grid_rowconfigure(0, weight=1)
        job_frame.grid_columnconfigure(0, weight=1)

        self.job_text = scrolledtext.ScrolledText(
            job_frame,
            wrap=tk.WORD,
            font=("Consolas", 10),
            bg="white",
            relief="solid",
            borderwidth=1,
        )
        self.job_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        # Score section
        score_frame = ttk.LabelFrame(
            content_frame, text="📊 Resume Score", padding="15"
        )
        score_frame.grid(row=2, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=15)
        score_frame.grid_columnconfigure(0, weight=1)

        self.score_label = ttk.Label(
            score_frame,
            text="Upload resume and job description, then click 'Score Resume' to see results",
            style="Score.TLabel",
            justify=tk.CENTER,
        )
        self.score_label.grid(row=0, column=0, sticky=(tk.W, tk.E))

        # Action buttons
        action_frame = ttk.Frame(main_frame)
        action_frame.grid(row=3, column=0, columnspan=2, pady=(15, 0))
        action_frame.grid_columnconfigure(0, weight=1)
        action_frame.grid_columnconfigure(1, weight=1)

        ttk.Button(
            action_frame,
            text="🎯 Score Resume",
            command=self.score_resume,
            style="Action.TButton",
        ).grid(row=0, column=0, padx=(0, 10), sticky=(tk.W, tk.E))

        ttk.Button(
            action_frame,
            text="✨ Rewrite Resume",
            command=self.rewrite_resume,
            style="Action.TButton",
        ).grid(row=0, column=1, padx=(10, 0), sticky=(tk.W, tk.E))

    def upload_resume(self):
        file_path = filedialog.askopenfilename(
            title="Select Resume File",
            filetypes=[("Document Files", "*.pdf *.docx *.txt"), ("All Files", "*.*")],
        )
        if file_path:
            try:
                content = parse_file(file_path)
                self.resume_text.delete(1.0, tk.END)
                self.resume_text.insert(tk.END, content)
                messagebox.showinfo("Success", "Resume uploaded successfully!")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to parse resume:\n{e}")

    def upload_job(self):
        file_path = filedialog.askopenfilename(
            title="Select Job Description File",
            filetypes=[("Document Files", "*.pdf *.docx *.txt"), ("All Files", "*.*")],
        )
        if file_path:
            try:
                content = parse_file(file_path)
                self.job_text.delete(1.0, tk.END)
                self.job_text.insert(tk.END, content)
                messagebox.showinfo("Success", "Job description uploaded successfully!")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to parse job description:\n{e}")

    def score_resume(self):
        resume = self.resume_text.get("1.0", tk.END)
        jd = self.job_text.get("1.0", tk.END)

        if not resume.strip() or not jd.strip():
            messagebox.showwarning(
                "Input Missing", "Please upload both resume and job description first."
            )
            return

        try:
            # Show loading message
            self.score_label.config(text="🔄 Calculating scores...")
            self.root.update()

            raw_scores = scorer.get_score(resume, jd)
            final_scores = scorer.calculate_final_score(raw_scores)

            display = (
                f"✅ Keyword Match: {final_scores['keyword_match']}%    "
                f"🧩 Section Match: {final_scores['section_match']}%\n"
                f"📈 Keyword Density: {final_scores['keyword_density']}%    "
                f"💯 Overall Score: {final_scores['overall_score']}%"
            )
            self.score_label.config(text=display)
        except Exception as e:
            messagebox.showerror("Scoring Error", f"Failed to calculate scores:\n{e}")
            self.score_label.config(text="❌ Error calculating scores")

    def rewrite_resume(self):
        resume = self.resume_text.get("1.0", tk.END)
        jd = self.job_text.get("1.0", tk.END)

        if not resume.strip() or not jd.strip():
            messagebox.showwarning(
                "Input Missing", "Please upload both resume and job description first."
            )
            return

        # Show processing message
        processing_dialog = tk.Toplevel(self.root)
        processing_dialog.title("Processing")
        processing_dialog.geometry("300x100")
        processing_dialog.configure(bg="#f0f0f0")
        tk.Label(
            processing_dialog,
            text="🔄 Rewriting resume...",
            font=("Arial", 12),
            bg="#f0f0f0",
        ).pack(expand=True)
        processing_dialog.transient(self.root)
        processing_dialog.grab_set()

        def close_processing():
            processing_dialog.destroy()

        threading.Thread(
            target=lambda: [self.launch_rewrite_window(resume, jd), close_processing()],
            daemon=True,
        ).start()

    def copy_all_to_clipboard(self, text_widget):
        text = text_widget.get("1.0", tk.END)
        self.root.clipboard_clear()
        self.root.clipboard_append(text.strip())
        messagebox.showinfo("Success", "Content copied to clipboard!")

    def export_to_docx(self, text):
        # Remove control characters and null bytes
        clean_text = re.sub(r"[\x00-\x1F\x7F]", "", text)

        filename = asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Document", "*.docx")],
            title="Save Resume As",
        )
        if filename:
            try:
                doc = Document()
                for line in clean_text.strip().split("\n"):
                    if line.strip():
                        doc.add_paragraph(line)
                doc.save(filename)
                messagebox.showinfo("Success", f"Resume exported to:\n{filename}")
            except Exception as e:
                messagebox.showerror("Export Failed", f"Error: {str(e)}")

    def launch_rewrite_window(self, resume, jd):
        rewrite_window = tk.Toplevel(self.root)
        rewrite_window.title("✨ Resume Rewriter Results")
        rewrite_window.geometry("900x700")
        rewrite_window.configure(bg="#f0f0f0")

        main_frame = ttk.Frame(rewrite_window, padding="15")
        main_frame.pack(fill=tk.BOTH, expand=True)

        title_label = ttk.Label(
            main_frame, text="📝 Rewritten Resume", font=("Arial", 14, "bold")
        )
        title_label.pack(pady=(0, 15))

        text_frame = ttk.Frame(main_frame)
        text_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 15))

        rewritten_text_widget = tk.Text(
            text_frame,
            wrap=tk.WORD,
            font=("Consolas", 10),
            bg="white",
            relief="solid",
            borderwidth=1,
        )
        scrollbar = ttk.Scrollbar(
            text_frame, orient=tk.VERTICAL, command=rewritten_text_widget.yview
        )
        rewritten_text_widget.configure(yscrollcommand=scrollbar.set)

        rewritten_text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # Insert initial loading message
        rewritten_text_widget.insert(tk.END, "⏳ Rewriting resume, please wait...\n")

        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X)

        ttk.Button(
            button_frame,
            text="📋 Copy All to Clipboard",
            command=lambda: self.copy_all_to_clipboard(rewritten_text_widget),
        ).pack(side=tk.LEFT, padx=(0, 10))

        ttk.Button(
            button_frame,
            text="💾 Export to DOCX",
            command=lambda: self.export_to_docx(
                rewritten_text_widget.get("1.0", tk.END)
            ),
        ).pack(side=tk.LEFT)

        def worker():
            def stream_update(section, rewritten):
                rewritten_text_widget.insert(
                    tk.END, f"## {section.capitalize()}\n{rewritten}\n\n"
                )
                rewritten_text_widget.see(tk.END)
                rewritten_text_widget.update_idletasks()

            try:
                for section, rewritten in rewrite_sections_streaming(resume, jd):
                    if section == "Error":

                        def update_error():
                            rewritten_text_widget.delete("1.0", tk.END)
                            rewritten_text_widget.insert(tk.END, f"⚠️ {rewritten}")

                        self.root.after(0, update_error)
                        return

                    # Call GUI updates from main thread
                    self.root.after(0, stream_update, section, rewritten)

            except Exception as e:

                def update_exception():
                    rewritten_text_widget.delete("1.0", tk.END)
                    rewritten_text_widget.insert(
                        tk.END, f"⚠️ Error rewriting resume:\n{str(e)}"
                    )

                self.root.after(0, update_exception)

        threading.Thread(target=worker, daemon=True).start()
